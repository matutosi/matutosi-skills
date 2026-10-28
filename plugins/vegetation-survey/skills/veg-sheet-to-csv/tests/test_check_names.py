#!/usr/bin/env python3
"""check_names.py の自己テスト (索引は手で作った小さなもの．xlsx の読み込みは openpyxl があるときだけ).

    python3 tests/test_check_names.py
"""

import json
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "scripts"))

import check_names as c  # noqa: E402

IDX = {
    "v": c.INDEX_VERSION,
    "hub": {
        "ツブラジイ": {"family": "Fagaceae", "family_jp": "ブナ", "status": "確定"},
        "ナナミノキ": {"family": "Aquifoliaceae", "family_jp": "モチノキ", "status": "確定"},
        "ベニシダ": {"family": "Dryopteridaceae", "family_jp": "オシダ", "status": "確定"},
        "ヒカゲア/ヒカゲイ": {"family": "X", "family_jp": "Y", "status": "確定"},
    },
    "alias": {
        "ツブラジイ": ["ツブラジイ"],
        "コジイ": ["ツブラジイ"],
        "ナナミノキ": ["ナナミノキ"],
        "ナナメノキ": ["ナナミノキ"],
        "ベニシダ": ["ベニシダ"],
        "カンムリベニシダ": ["ベニシダ"],
        "ヒカゲア": ["ヒカゲア/ヒカゲイ"],
        "ユレ": ["ツブラジイ", "ナナミノキ"],
        "ヤブコウジ": ["ベニシダ"],
    },
    "latin": {
        "ツブラジイ": {"GL": "Castanopsis cuspidata", "YL": "Castanopsis cuspidata var. cuspidata"},
        "ナナミノキ": {"GL": "Ilex chinensis"},
        "ベニシダ": {"WF": "Dryopteris erythrosora"},
    },
}


def mk(names):
    return {"site": {}, "species": [{"階層": "T1", "種名": n} for n in names]}


class CheckNameTest(unittest.TestCase):
    def test_status(self):
        cases = {
            "ツブラジイ": "ok",
            "コジイ": "alias",
            "ｺｼﾞｲ": "alias",      # 半角カナも正規化して照合
            "こじい": "alias",    # ひらがなも
            "ナナメノキ": "alias",
            "ヒカゲア": "ambiguous",
            "ユレ": "ambiguous",
            "ヤブコウジー": "similar",
            "ゼンゼンチガウ": "none",
            "スゲ属": "skip",
            "ササ sp.": "skip",
            "コジイ?": "skip",
            "Carex": "skip",
            "": "skip",
        }
        for name, want in cases.items():
            self.assertEqual(c.check_name(IDX, name)["status"], want, name)

    def test_alias_result(self):
        r = c.check_name(IDX, "コジイ")
        self.assertEqual((r["hub"], r["family"]), ("ツブラジイ", "ブナ"))

    def test_latin_priority(self):
        self.assertEqual(c.check_name(IDX, "ツブラジイ")["latin"], "Castanopsis cuspidata var. cuspidata")
        self.assertIn("違う", c.check_name(IDX, "ツブラジイ")["note"])
        self.assertEqual(c.check_name(IDX, "ツブラジイ", ["GL"])["latin"], "Castanopsis cuspidata")
        self.assertEqual(c.check_name(IDX, "ベニシダ")["latin"], "Dryopteris erythrosora")

    def test_ambiguous_candidates(self):
        self.assertEqual(c.check_name(IDX, "ヒカゲア")["candidates"], ["ヒカゲア", "ヒカゲイ"])


class ProcessTest(unittest.TestCase):
    def test_keep_and_fill(self):
        d = mk(["コジイ", "ナナミノキ"])
        d["species"][1]["学名"] = "Ilex 手書き"
        c.process(d, IDX, "keep", ("latin", "family"))
        self.assertEqual(d["species"][0]["種名"], "コジイ")           # 置き換えない
        self.assertEqual(d["species"][0]["学名"], "Castanopsis cuspidata var. cuspidata")
        self.assertEqual(d["species"][0]["科名"], "ブナ")
        self.assertEqual(d["species"][1]["学名"], "Ilex 手書き")      # 既存の値は上書きしない

    def test_rename_hub(self):
        d = mk(["コジイ", "ヒカゲア", "ゼンゼンチガウ"])
        c.process(d, IDX, "hub")
        self.assertEqual([s["種名"] for s in d["species"]], ["ツブラジイ", "ヒカゲア", "ゼンゼンチガウ"])

    def test_no_fill_by_default(self):
        d = mk(["コジイ"])
        c.process(d, IDX)
        self.assertNotIn("学名", d["species"][0])

    def test_sheets(self):
        d = {"sheets": [mk(["コジイ"]), mk(["コジイ", "ベニシダ"])]}
        rows, cache = c.process(d, IDX, "hub", ("latin",))
        self.assertEqual(len(rows), 3)
        self.assertEqual(len(cache), 2)
        self.assertTrue(all(s["species"][0]["種名"] == "ツブラジイ" for s in d["sheets"]))

    def test_report_order(self):
        _, cache = c.process(mk(["ツブラジイ", "ゼンゼンチガウ", "ヒカゲア"]), IDX)
        lines = c.format_report(cache).splitlines()
        self.assertEqual([l.split("\t")[0] for l in lines[1:]], ["ambiguous", "none", "ok"])


class XlsxTest(unittest.TestCase):
    def test_build_index_from_xlsx(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl が無い")
        wb = openpyxl.Workbook()
        h = wb.active
        h.title = "Hub_data"
        h.append(["all_name", "Hub name", "lato/stricto", "Family ID", "Family name", "Family name (JP)", "GL", "SF", "WF", "YL", "status", "message"])
        h.append(["ツブラジイ", "ツブラジイ", None, 1, "Fagaceae", "ブナ", "GL_1", "", "", "YL_1", "確定", None])
        h.append(["コジイ", "ツブラジイ", None, 1, "Fagaceae", "ブナ", "GL_1", "", "", "YL_1", "確定", None])
        j = wb.create_sheet("JN_dataset")
        j.append(["ID", "Family ID", "Family name", "Family name (JP)", "common name", "another name", "another name ID", "note 1", "note 2", "scientific name with author", "scientific name without author"])
        j.append(["GL_1", 1, "Fagaceae", "ブナ", "ツブラジイ", "ツブラジイ", 0, None, None, "Castanopsis cuspidata (Thunb.) Schottky", "Castanopsis cuspidata"])
        path = os.path.join(tempfile.mkdtemp(), "w.xlsx")
        wb.save(path)
        idx = c.load_index(path)
        self.assertTrue(os.path.exists(path + ".index.json"))
        r = c.check_name(idx, "コジイ")
        self.assertEqual((r["status"], r["hub"], r["latin"], r["family"]), ("alias", "ツブラジイ", "Castanopsis cuspidata", "ブナ"))
        # 2回目は JSON の索引から読む
        self.assertEqual(c.load_index(path), idx)

        data = os.path.join(os.path.dirname(path), "d.json")
        with open(data, "w", encoding="utf-8") as f:
            json.dump(mk(["コジイ"]), f, ensure_ascii=False)
        out = os.path.join(os.path.dirname(path), "o.json")
        rc = c.main([data, "--checklist", path, "--fill", "latin", "--rename", "hub", "-o", out, "--strict"])
        self.assertEqual(rc, 0)
        with open(out, encoding="utf-8") as f:
            sp = json.load(f)["species"][0]
        self.assertEqual((sp["種名"], sp["学名"]), ("ツブラジイ", "Castanopsis cuspidata"))

    def test_strict_exit_code(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl が無い")
        path = os.path.join(tempfile.mkdtemp(), "w.xlsx")
        import openpyxl as ox
        wb = ox.Workbook()
        wb.active.title = "Hub_data"
        wb.active.append(["all_name", "Hub name"])
        wb.create_sheet("JN_dataset").append(["ID"])
        wb.save(path)
        data = os.path.join(os.path.dirname(path), "d.json")
        with open(data, "w", encoding="utf-8") as f:
            json.dump(mk(["ゼンゼンチガウ"]), f, ensure_ascii=False)
        self.assertEqual(c.main([data, "--checklist", path, "--strict"]), 2)
        self.assertEqual(c.main([data, "--checklist", path]), 0)


if __name__ == "__main__":
    unittest.main(verbosity=1)
