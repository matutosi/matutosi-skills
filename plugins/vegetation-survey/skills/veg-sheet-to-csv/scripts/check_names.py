#!/usr/bin/env python3
"""読み取った種名を「維管束植物和名チェックリスト」(GBIF Japan) と照合する.

    python check_names.py data.json [-o checked.json] [--rename keep|hub]
                          [--fill latin,family] [--source YL,GL,WF,SF]
                          [--checklist FILE.xlsx | --download] [--report out.tsv] [--strict]

build_csv.py に渡す JSON (sheets 形式でも 1 枚形式でも可) の species[].種名 を照合し，
  ok         標準和名 (Hub name) と一致
  alias      異名と一致 (標準和名を示す．--rename hub で置き換える)
  ambiguous  和名の指す範囲が資料で異なる・複数の標準和名に当たる (人が決める)
  similar    一致しないが似た名前がある (候補を示す．自動では直さない)
  none       見つからない
  skip       属・科以上，sp.，cf.，? を含む，カナ以外 (照合しない)
を報告する．--fill で，空の「学名」「科名」だけをチェックリストから埋める (既存の値は上書きしない)．

チェックリスト: 山ノ内崇志ほか (2019) 維管束植物和名チェックリスト ver. 1.10 (CC BY 4.0)
  https://gbif.jp/activities/checklist/wamei_checklist_110/
xlsx は --download で取得して ~/.cache/veg-sheet-to-csv/ に置く (環境変数 VEG_CHECKLIST でも指定できる)．
xlsx を読むには openpyxl が要る (読んだ索引は xlsx の隣に JSON で残し，2回目から openpyxl は要らない)．
"""

import argparse
import difflib
import json
import os
import sys
import unicodedata
import urllib.request

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

URL = "https://gbif.jp/activities/checklist/wamei_checklist_110/excel/wamei_checklist_ver.1.10.xlsx"
CACHE_DIR = os.path.join(os.path.expanduser("~"), ".cache", "veg-sheet-to-csv")
DEFAULT_XLSX = os.path.join(CACHE_DIR, "wamei_checklist_ver.1.10.xlsx")
SOURCES = ["YL", "GL", "WF", "SF"]  # 学名の優先順 (YList・Green List・野生植物・シダ標準図鑑)
INDEX_VERSION = 1


def norm(s):
    """全角カナ・空白なしに寄せる."""
    s = unicodedata.normalize("NFKC", str(s or "")).strip()
    s = "".join(chr(ord(c) + 0x60) if "ぁ" <= c <= "ゖ" else c for c in s)
    return s.replace(" ", "").replace("　", "")


def build_index(xlsx):
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl が要る: pip install openpyxl")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    hub, alias, latin = {}, {}, {}
    for r in wb["Hub_data"].iter_rows(min_row=2, values_only=True):
        if not r[0] or not r[1]:
            continue
        name, h = norm(r[0]), norm(r[1])
        hub.setdefault(h, {"family": r[4] or "", "family_jp": r[5] or "", "status": r[10] or ""})
        lst = alias.setdefault(name, [])
        if h not in lst:
            lst.append(h)
    for r in wb["JN_dataset"].iter_rows(min_row=2, values_only=True):
        if not r[4] or not (r[10] or r[9]):
            continue
        src = str(r[0] or "")[:2]
        sci = str(r[10] or r[9]).strip()
        latin.setdefault(norm(r[4]), {}).setdefault(src, sci)
    return {"v": INDEX_VERSION, "hub": hub, "alias": alias, "latin": latin}


def load_index(xlsx):
    cache = xlsx + ".index.json"
    if os.path.exists(cache) and os.path.getmtime(cache) >= os.path.getmtime(xlsx):
        with open(cache, encoding="utf-8") as f:
            idx = json.load(f)
        if idx.get("v") == INDEX_VERSION:
            return idx
    idx = build_index(xlsx)
    with open(cache, "w", encoding="utf-8") as f:
        json.dump(idx, f, ensure_ascii=False)
    return idx


def is_checkable(name):
    if not name:
        return False
    if any(m in name for m in ("sp.", "spp.", "cf.", "?", "？", "不明", "Unknown")):
        return False
    if name.endswith(("属", "科", "類", "sp")):
        return False
    return all("ァ" <= c <= "ヺ" or c in "ー・ヽヾ" for c in name)


def pick_latin(idx, hub_name, sources):
    """標準和名から学名を選ぶ。(学名, 資料間で食い違うか) を返す."""
    by_src = idx["latin"].get(hub_name, {})
    for s in sources:
        if s in by_src:
            return by_src[s], len(set(by_src.values())) > 1
    return "", False


def check_name(idx, raw, sources=SOURCES):
    """1つの種名を照合する。dict(status, hub, candidates, latin, family, note) を返す."""
    name = norm(raw)
    res = {"input": raw, "status": "none", "hub": "", "candidates": [], "latin": "", "family": "", "note": ""}
    if not is_checkable(name):
        res["status"] = "skip"
        return res
    hubs = idx["hub"]
    aliases = idx["alias"].get(name, [])
    if name in hubs:
        hub = name
        res["status"] = "ok"
        others = [h for h in aliases if h != name]
        if others:
            res["note"] = "同じ名前が別の標準和名の異名でもある: " + "，".join(others)
    elif len(aliases) == 1:
        hub = aliases[0]
        res["status"] = "alias"
    elif len(aliases) > 1:
        res.update(status="ambiguous", candidates=aliases, note="複数の標準和名に当たる")
        return res
    else:
        res["candidates"] = difflib.get_close_matches(name, list(idx["alias"]), n=3, cutoff=0.7)
        if res["candidates"]:
            res["status"] = "similar"
        return res
    if "/" in hub:
        res.update(status="ambiguous", candidates=hub.split("/"), note="資料により和名の指す範囲が違う")
        return res
    res["hub"] = hub
    res["family"] = hubs[hub].get("family_jp", "")
    res["latin"], differ = pick_latin(idx, hub, sources)
    if differ:
        res["note"] = (res["note"] + "；" if res["note"] else "") + "学名が資料間で違う"
    return res


def iter_sheets(data):
    return data["sheets"] if "sheets" in data else [data]


def process(data, idx, rename="keep", fill=(), sources=SOURCES):
    """data を直接書き換え、(全結果の一覧, 名前ごとの結果) を返す."""
    cache, rows = {}, []
    for si, sheet in enumerate(iter_sheets(data), 1):
        for sp in sheet["species"]:
            raw = str(sp.get("種名", ""))
            if raw not in cache:
                cache[raw] = check_name(idx, raw, sources)
            r = cache[raw]
            rows.append((si, sp.get("階層", ""), r))
            if r["status"] == "alias" and rename == "hub":
                sp["種名"] = r["hub"]
            if r["hub"] and r["status"] in ("ok", "alias"):
                if "latin" in fill and not str(sp.get("学名", "")).strip() and r["latin"]:
                    sp["学名"] = r["latin"]
                if "family" in fill and not str(sp.get("科名", "")).strip() and r["family"]:
                    sp["科名"] = r["family"]
    return rows, cache


def format_report(cache):
    lines = ["状態\t入力\t標準和名\t学名\t科名\t候補\t備考"]
    order = {"ambiguous": 0, "none": 1, "similar": 2, "alias": 3, "ok": 4, "skip": 5}
    for r in sorted(cache.values(), key=lambda r: (order[r["status"]], r["input"])):
        lines.append("\t".join([r["status"], r["input"], r["hub"], r["latin"], r["family"], "，".join(r["candidates"]), r["note"]]))
    return "\n".join(lines)


def resolve_xlsx(a):
    path = a.checklist or os.environ.get("VEG_CHECKLIST") or DEFAULT_XLSX
    if not os.path.exists(path):
        if not a.download:
            sys.exit("チェックリストが無い: %s\n--download で取得するか，--checklist で xlsx を指定する (出典: %s)" % (path, URL))
        os.makedirs(os.path.dirname(path), exist_ok=True)
        print("ダウンロード:", URL, file=sys.stderr)
        urllib.request.urlretrieve(URL, path)
    return path


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json")
    ap.add_argument("-o", "--out", help="照合・補完したあとの JSON (省略すると書かない)")
    ap.add_argument("--rename", choices=["keep", "hub"], default="keep", help="異名を標準和名に置き換える (既定は置き換えない)")
    ap.add_argument("--fill", default="", help="空の列だけ埋める: latin,family")
    ap.add_argument("--source", default=",".join(SOURCES), help="学名の優先順")
    ap.add_argument("--checklist", help="チェックリストの xlsx")
    ap.add_argument("--download", action="store_true", help="xlsx が無ければ取得する")
    ap.add_argument("--report", help="報告を TSV で書く")
    ap.add_argument("--strict", action="store_true", help="ambiguous・none・similar があれば終了コード 2")
    a = ap.parse_args(argv)

    idx = load_index(resolve_xlsx(a))
    with open(a.json, encoding="utf-8") as f:
        data = json.load(f)
    fill = tuple(x for x in a.fill.split(",") if x)
    rows, cache = process(data, idx, a.rename, fill, a.source.split(","))
    rep = format_report(cache)
    print(rep)
    counts = {}
    for r in cache.values():
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    print("\n集計:", "，".join("%s %d" % kv for kv in sorted(counts.items())), file=sys.stderr)
    if a.report:
        with open(a.report, "w", encoding="utf-8-sig", newline="") as f:
            f.write(rep + "\n")
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=1)
    bad = sum(counts.get(k, 0) for k in ("ambiguous", "none", "similar"))
    return 2 if a.strict and bad else 0


if __name__ == "__main__":
    sys.exit(main())
